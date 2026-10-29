"""Order ingestion: Shriram's state-wise LAN sheets (Google Sheet / .xlsx / .csv) in Orders Incoming.

Per row: LAN format, state (column → tab name → file name), stamp duty (column → state default),
state active, amount allowed for the state, duplicates in file, duplicates in the ledger.

Same LAN arriving for a DIFFERENT state while still unfulfilled → both copies go to CONFLICT and nobody
gets a paper until ops picks the right state (Order Actions tab → KEEP_STATE).

A file is refused whole (E201, → Orders Rejected) rather than half-read when a tab has more filled rows than
orders.max_rows_per_tab, or an .xlsx column holds formulas saved without their results (they would read as blank).
"""
from __future__ import annotations

import csv
import io
import itertools
import time
from datetime import datetime, timezone

from .constants import ORDERS_INCOMING, ORDERS_PROCESSED, ORDERS_REJECTED
from .context import Ctx
from .errors import line
from .extract import ParsedOrderFile, parse_order_tabs
from .google_clients import CSV, GSHEET, XLS, XLSX, parse_time


# L-2: Drive items seen without a modifiedTime → first time this process saw them (in-memory: a restart just waits again).
_first_seen: dict[str, datetime] = {}


def _modified(ctx: Ctx, f: dict) -> datetime:
    """Drive's modifiedTime, or – when an item lacks it – when we first saw it, so it still waits out the settle window."""
    if f.get("modifiedTime"):
        return parse_time(f["modifiedTime"])
    if f["id"] not in _first_seen:
        _first_seen[f["id"]] = datetime.now(timezone.utc)
        ctx.event("INFO", "Orders", f"'{f['name']}' has no modifiedTime in Drive – waiting the settle window from now",
                  ref=f["id"])
    return _first_seen[f["id"]]


def run_ingest(ctx: Ctx, deadline: float) -> None:
    g, s = ctx.google, ctx.settings
    settle = int(s.get("orders.settle_minutes", 15)) * 60
    for f in g.list_children(s.folder(ORDERS_INCOMING), folders=False):
        if time.time() > deadline:
            return
        age = (datetime.now(timezone.utc) - _modified(ctx, f)).total_seconds()
        if age < settle:
            continue
        try:
            _ingest_file(ctx, f)
        except Exception as e:
            ctx.stats.files_rejected += 1
            ctx.event("ERROR", "Orders", line("E201", f"'{f['name']}' rejected: {e}"), ref=f["id"])
            if not ctx.dry:
                g.move_rename(f["id"], s.folder(ORDERS_REJECTED))


def _cell(v) -> str:
    return "" if v is None else (str(int(v)) if isinstance(v, float) and v.is_integer() else str(v))


def _capped(title: str, rows, cap: int) -> list[list[str]]:
    """Rows of one tab, refusing more than `cap` filled rows. Blank rows don't count; trailing ones are dropped
    (interior ones are kept so row numbers in messages stay right)."""
    out, blanks, filled = [], 0, 0
    for r in rows:
        if not any(str(c).strip() for c in r):
            blanks += 1
            continue
        filled += 1
        if filled > cap:
            raise ValueError(f"tab '{title}' has more than {cap} rows (orders.max_rows_per_tab) – split the file "
                             f"or raise the cap in config.yaml")
        out.extend([[] for _ in range(blanks)])
        blanks = 0
        out.append(r)
    return out


def _xlsx_tabs(data: bytes, cap: int) -> list[tuple[str, list[list[str]]]]:
    """Cached values (data_only) read alongside the formulas (a second streaming pass). A column whose formulas
    ALL lack a cached value was saved by a tool that never calculated them (LibreOffice headless, some exporters):
    every such cell reads None, so the LAN/state/amount would look blank – refuse the file with the real reason."""
    from openpyxl import load_workbook
    values = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    out = []
    for ws, wf in zip(values.worksheets, formulas.worksheets):
        uncached: dict[int, int] = {}   # column → formula cells with no cached value
        cached: dict[int, int] = {}     # column → formula cells with one
        def rows():
            for rv, rf in itertools.zip_longest(ws.iter_rows(values_only=True), wf.iter_rows(values_only=True),
                                                fillvalue=()):
                for i, fv in enumerate(rf):
                    if isinstance(fv, str) and fv.startswith("="):
                        hit = uncached if (rv[i] if i < len(rv) else None) is None else cached
                        hit[i] = hit.get(i, 0) + 1
                yield [_cell(v) for v in rv]
        tab = _capped(ws.title, rows(), cap)
        bad = sorted(c for c in uncached if not cached.get(c))
        if bad:
            cols = ", ".join(_col_letter(c) for c in bad)
            raise ValueError(f"tab '{ws.title}' column(s) {cols}: file has formulas with no cached values – open in "
                             f"Excel and re-save (or paste as values), then upload again")
        out.append((ws.title, tab))
    return out


def _col_letter(i: int) -> str:
    from openpyxl.utils import get_column_letter
    return get_column_letter(i + 1)


def _read_tabs(ctx: Ctx, f: dict) -> list[tuple[str, list[list[str]]]]:
    g, mt = ctx.google, f.get("mimeType")
    cap = int(ctx.settings.get("orders.max_rows_per_tab", 50000))
    if mt == GSHEET:
        return [(t, _capped(t, rows, cap)) for t, rows in g.read_tabs(f["id"])]
    data = g.download(f["id"])
    if mt == CSV or f["name"].lower().endswith(".csv"):
        return [(f["name"], _capped(f["name"], csv.reader(io.StringIO(data.decode("utf-8-sig", "replace"))), cap))]
    if mt in (XLSX, XLS) or f["name"].lower().endswith((".xlsx", ".xls")):
        return _xlsx_tabs(data, cap)
    raise ValueError(f"unsupported file type {mt} – send a Google Sheet, .xlsx or .csv")


def _ingest_file(ctx: Ctx, f: dict) -> None:
    g, s, db = ctx.google, ctx.settings, ctx.db
    parsed: ParsedOrderFile = parse_order_tabs(_read_tabs(ctx, f), f["name"], s, ctx.resolver)
    if not parsed.header_found:
        raise ValueError("no header row with a LAN column in any tab (check orders.columns in config.yaml)")
    if parsed.invalid and s.get("orders.reject_file_on_any_invalid_row", False):
        raise ValueError(f"{len(parsed.invalid)} invalid row(s): " + " | ".join(parsed.invalid[:15]))

    existing = db.existing_orders([o.lan_key for o in parsed.valid])
    to_insert, dups, conflicts = [], list(parsed.duplicates), []
    for o in parsed.valid:
        prev = existing.get(o.lan_key, [])
        base = (o.lan, o.lan_key, o.state, o.denomination, f["name"], f["id"], o.tab, o.row_no)
        if not prev:
            to_insert.append(base + ("PENDING", None))
        elif all(p["state"] == o.state for p in prev):
            dups.append(f"{o.tab}!row {o.row_no}: {o.lan} already received (status {prev[-1]['status']})")
        elif any(p["status"] in ("DONE", "ALLOCATING") for p in prev):
            dups.append(f"{o.tab}!row {o.row_no}: {o.lan} for {o.state} was ALREADY FULFILLED for "
                        f"{[p['state'] for p in prev]} – not processed")
        else:
            note = f"LAN received for several states: {sorted({p['state'] for p in prev} | {o.state})} – pick one"
            if not ctx.dry:
                db.mark_conflict([p["id"] for p in prev], note)
            to_insert.append(base + ("CONFLICT", note))
            conflicts.append(f"{o.lan} ({note})")

    if not to_insert and not parsed.valid:
        raise ValueError(f"no valid rows ({len(parsed.invalid)} invalid, {len(dups)} duplicates): "
                         + " | ".join((parsed.invalid + dups)[:15]))

    if ctx.dry:
        for r in to_insert:
            ctx.preview.append(["ORDER", f["name"], f"{r[6]}!row {r[7]}", r[0], r[2], r[3], r[8], r[9] or "OK"])
            if r[8] == "PENDING":
                ctx.dryPending.append({"LAN ID": r[0], "State": r[2], "Denomination": r[3], "Source File": r[4],
                                       "Source Tab": r[6], "Source Row": r[7]})
        for m in parsed.invalid + dups:
            ctx.preview.append(["ORDER", f["name"], "", "", "", "", "SKIPPED", m])
    else:
        db.insert_orders(to_insert)
        g.move_rename(f["id"], s.folder(ORDERS_PROCESSED))

    ctx.stats.orders_in += sum(1 for r in to_insert if r[8] == "PENDING")
    ctx.stats.rows_rejected += len(parsed.invalid)
    ctx.stats.duplicates += len(dups)
    ctx.stats.conflicts += len(conflicts)
    ctx.event("INFO", "Orders", f"'{f['name']}': {sum(1 for r in to_insert if r[8] == 'PENDING')} LAN(s) queued",
              ref=f["id"])
    if parsed.invalid:
        ctx.event("WARN", "Orders", line("E203", f"'{f['name']}': {len(parsed.invalid)} row(s) skipped – "
                  + " | ".join(parsed.invalid[:20])))
    if dups:
        ctx.event("WARN", "Orders", line("E205", f"'{f['name']}': {len(dups)} duplicate(s) – "
                  + " | ".join(dups[:20])))
    if conflicts:
        ctx.event("WARN", "Orders", line("E206", f"'{f['name']}': {len(conflicts)} LAN(s) in CONFLICT – resolve in "
                  "Order Actions: " + " | ".join(conflicts[:20])))
