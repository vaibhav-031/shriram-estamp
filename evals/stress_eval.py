"""Stress-test eval loop for the live e-stamp pipeline.

Flow (see evals/DESIGN.md):
  reset → build assets from real templates → upload to Drive → drive cycles → verify vs oracle → report.

The oracle is the `LAN-Stamp Mapping` tab in each order workbook: it maps every LAN to its exact
certificate, stamp duty, source PDF and page. We assert the pipeline's output against that ground truth,
not against memorized counts.

Runs against REAL Drive / OCR / Postgres (config.yaml). `--live` issues real certificates – only use on a
pre-production environment. `--reset` TRUNCATEs the schema first (same as `estamp reset-test-data`).

  python evals/stress_eval.py --scenario both --reset --live
  python evals/stress_eval.py --scenario inbox --states KA,BR --papers 5
"""
from __future__ import annotations

import argparse
import io
import sys
import time
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader, PdfWriter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from estamp.context import Ctx  # noqa: E402
from estamp.extract import StateResolver  # noqa: E402
from estamp.settings import load_settings  # noqa: E402

TEMPLATES = ROOT / "evals" / "templates"
PREFIX = "EVAL"   # every asset we create is named EVAL_* so a human can spot test data at a glance


# ----------------------------------------------------------------- templates → papers (the oracle)

@dataclass
class Paper:
    lan: str
    state: str          # resolved code, e.g. "KA"
    state_raw: str
    cert: str
    amount: int
    pdf_file: str
    page: int           # 1-based page in pdf_file


def _col(header: list[str], *names: str) -> int:
    low = [str(h or "").strip().lower() for h in header]
    for n in names:
        if n.lower() in low:
            return low.index(n.lower())
    raise KeyError(f"none of {names} in {header}")


def load_papers(resolver: StateResolver) -> list[Paper]:
    """Every LAN→paper mapping across all template workbooks, grouped later by state."""
    papers: list[Paper] = []
    for xlsx in sorted((TEMPLATES / "orders").glob("*.xlsx")):
        wb = load_workbook(xlsx, data_only=True)
        if "LAN-Stamp Mapping" not in wb.sheetnames:
            continue
        ws = wb["LAN-Stamp Mapping"]
        rows = list(ws.iter_rows(values_only=True))
        hdr = [str(c or "") for c in rows[0]]
        ci = {k: _col(hdr, *v) for k, v in {
            "lan": ("LOAN NO",), "state": ("STATE",), "cert": ("Certificate / Serial No.",),
            "amount": ("Stamp Duty",), "pdf": ("Sample PDF",), "page": ("Page",)}.items()}
        for r in rows[1:]:
            if not r or not r[ci["lan"]]:
                continue
            state_raw = str(r[ci["state"]])
            papers.append(Paper(
                lan=str(r[ci["lan"]]).strip(), state=resolver.resolve(state_raw), state_raw=state_raw,
                cert=str(r[ci["cert"]]).strip(), amount=int(r[ci["amount"]]),
                pdf_file=str(r[ci["pdf"]]).strip(), page=int(r[ci["page"]])))
    return papers


# ----------------------------------------------------------------- scenario + oracle

@dataclass
class Group:
    state: str
    amount: int
    papers: list[Paper]          # chosen papers (uploaded)
    lans: list[str]              # chosen LANs (order sheet, scenario 1)
    active: bool
    single_amount: bool          # order sheet carries no duty column → amount must be inferable
    # oracle, filled by build_oracle
    valid_lans: int = 0
    issuable: int = 0
    pending: int = 0
    available: int = 0           # papers expected AVAILABLE after the run
    assigned: int = 0            # papers expected ASSIGNED (issued)
    review: int = 0              # papers expected in Needs Review (inactive-state papers)


@dataclass
class Scenario:
    scenario: str                # orders | inbox | both
    groups: list[Group]
    dup_upload: bool


def build_scenario(settings, papers: list[Paper], states: list[str] | None,
                   n_lans: int, n_papers: int, scenario: str, dup_upload: bool) -> Scenario:
    by_state: dict[str, list[Paper]] = {}
    for p in papers:
        by_state.setdefault(p.state, []).append(p)
    groups = []
    for state, ps in sorted(by_state.items()):
        if states and state not in states:
            continue
        ps = sorted(ps, key=lambda p: (p.pdf_file, p.page))
        st = settings.states.get(state)
        active = bool(st and st.active)
        single = bool(st and len(st.amounts) == 1)
        amount = ps[0].amount
        chosen_papers = ps[:n_papers]
        chosen_lans = [p.lan for p in ps[:n_lans]]
        groups.append(Group(state=state, amount=amount, papers=chosen_papers, lans=chosen_lans,
                            active=active, single_amount=single))
    build_oracle(groups, scenario)
    return Scenario(scenario=scenario, groups=groups, dup_upload=dup_upload)


def build_oracle(groups: list[Group], scenario: str) -> None:
    for g in groups:
        good = len(g.papers) if g.active else 0   # inactive-state papers can't be stocked → Needs Review
        g.review = len(g.papers) - good
        if scenario == "inbox":
            g.valid_lans = g.issuable = g.pending = g.assigned = 0
            g.available = good
        else:
            # An order sheet with no duty column needs the state active AND single-amount to parse a LAN as valid.
            g.valid_lans = len(g.lans) if (g.active and g.single_amount) else 0
            g.issuable = g.assigned = min(g.valid_lans, good)
            g.pending = g.valid_lans - g.issuable
            g.available = good - g.issuable   # good papers left over after allocation


# ----------------------------------------------------------------- asset upload (real Drive)

def _merged_pdf(papers: list[Paper], readers: dict[str, PdfReader]) -> bytes:
    w = PdfWriter()
    for p in papers:
        w.add_page(readers[p.pdf_file].pages[p.page - 1])
    out = io.BytesIO()
    w.write(out)
    return out.getvalue()


def upload_assets(google, settings, sc: Scenario) -> None:
    readers = {f.name: PdfReader(str(f)) for f in (TEMPLATES / "pdf").glob("*.pdf")}
    inbox = settings.folder("upload_inbox")
    incoming = settings.folder("orders_incoming")
    # Papers MUST sit inside a known vendor folder, never the inbox root (E101, GLOSSARY "Vendor folder").
    vendor = next((v.name for v in settings.vendors.values() if v.active), None)
    if vendor is None:
        raise RuntimeError("no active vendor configured – intake would reject every upload (E101)")
    vendor_folder = google.folder_path(inbox, [vendor])
    for g in sc.groups:
        if not g.papers:
            continue
        data = _merged_pdf(g.papers, readers)
        name = f"{PREFIX}_{g.state}_{g.amount}_x{len(g.papers)}.pdf"
        print(f"  upload paper PDF  → inbox/{vendor}: {name} ({len(g.papers)} pages)")
        google.upload_pdf(vendor_folder, name, data)
        if sc.dup_upload:
            print(f"  upload DUPLICATE  → inbox/{vendor}: {name} (same bytes, exact-copy test)")
            google.upload_pdf(vendor_folder, f"{PREFIX}_{g.state}_{g.amount}_DUP.pdf", data)
        if sc.scenario != "inbox" and g.lans:
            sheet = f"{PREFIX}_{g.state}_LANs"
            header = ["STATE", "LOAN NO", "Customer Name"]
            print(f"  create order sheet → incoming: {sheet} ({len(g.lans)} LANs)")
            sid = google.create_sheet(incoming, sheet, header)
            state_raw = g.papers[0].state_raw
            google.append_rows(sid, sheet, [[state_raw, lan, f"Eval {lan}"] for lan in g.lans])


# ----------------------------------------------------------------- drive cycles

def drive_cycles(settings, db, google, max_cycles: int) -> int:
    from estamp.cycle import run_cycle
    ran = 0
    for i in range(max_cycles):
        print(f"  cycle {i + 1}/{max_cycles} …")
        run_cycle(Ctx(settings, db, google))
        ran += 1
        pg = db.pending_groups()
        if not pg or all(g["avail"] == 0 for g in pg):
            break
    return ran


# ----------------------------------------------------------------- verify vs oracle

@dataclass
class Report:
    checks: list[tuple[bool, str]] = field(default_factory=list)

    def check(self, ok: bool, msg: str) -> None:
        self.checks.append((bool(ok), msg))

    @property
    def failures(self) -> list[str]:
        return [m for ok, m in self.checks if not ok]


def verify(db, settings, sc: Scenario) -> Report:
    rep = Report()
    certs_of = {g.state: {p.cert for p in g.papers} for g in sc.groups}

    for g in sc.groups:
        where = f"{g.state} ₹{g.amount}"
        # ---- papers stocked: AVAILABLE spare + ASSIGNED issuable, and no good paper lost
        avail = _n(db, "SELECT count(*) n FROM papers WHERE state=%s AND denomination=%s AND status='AVAILABLE'",
                   (g.state, g.amount))
        assigned = _n(db, "SELECT count(*) n FROM papers WHERE state=%s AND denomination=%s AND status='ASSIGNED'",
                      (g.state, g.amount))
        rep.check(avail == g.available, f"{where}: AVAILABLE papers {avail} == expected {g.available}")
        if sc.scenario != "inbox":
            rep.check(assigned == g.assigned, f"{where}: ASSIGNED papers {assigned} == issuable {g.assigned}")
        if g.review:
            rev = _n(db, "SELECT count(*) n FROM papers WHERE state=%s AND status='NEEDS_REVIEW'", (g.state,))
            rep.check(rev >= g.review, f"{where}: inactive-state papers in Needs Review {rev} >= {g.review}")
        # every stocked paper's cert belongs to this group's known certs
        stocked = {r["cert_no"] for r in db.q(
            "SELECT cert_no FROM papers WHERE state=%s AND denomination=%s AND status IN ('AVAILABLE','ASSIGNED')",
            (g.state, g.amount))}
        rep.check(stocked <= certs_of[g.state],
                  f"{where}: stocked certs ⊆ template certs (extras: {sorted(stocked - certs_of[g.state])[:3]})")

        if sc.scenario == "inbox":
            continue
        # ---- orders: valid_lans split into issuable DONE + pending PENDING
        done = _n(db, "SELECT count(*) n FROM orders WHERE state=%s AND denomination=%s AND status='DONE'",
                  (g.state, g.amount))
        pend = _n(db, "SELECT count(*) n FROM orders WHERE state=%s AND denomination=%s AND status='PENDING'",
                  (g.state, g.amount))
        rep.check(done == g.issuable, f"{where}: orders DONE {done} == issuable {g.issuable}")
        rep.check(pend == g.pending, f"{where}: orders PENDING {pend} == expected {g.pending}")

    if sc.scenario != "inbox":
        # ---- global allocation integrity (the money-critical invariants)
        allocs = db.q("SELECT * FROM allocations WHERE cancelled_at IS NULL")
        total_issuable = sum(g.issuable for g in sc.groups)
        rep.check(len(allocs) == total_issuable, f"allocations {len(allocs)} == total issuable {total_issuable}")
        rep.check(len({a["paper_id"] for a in allocs}) == len(allocs), "no paper issued twice")
        rep.check(len({a["order_id"] for a in allocs}) == len(allocs), "no LAN got two papers")
        all_certs = set().union(*certs_of.values()) if certs_of else set()
        rep.check(all(a["cert_no"] in all_certs for a in allocs), "every issued cert is a known template cert")
        rep.check(all(a["pdf_name"] == f"{a['lan']}.pdf" for a in allocs), "every issued file named <LAN>.pdf")
        mism = _n(db, "SELECT count(*) n FROM allocations WHERE verified_cert IS DISTINCT FROM cert_no "
                  "AND cancelled_at IS NULL")
        rep.check(mism == 0, f"second-OCR verified_cert matches ledger ({mism} mismatches)")
        stuck = _n(db, "SELECT count(*) n FROM orders WHERE status IN ('ALLOCATING','ERROR')")
        rep.check(stuck == 0, f"no order left stuck ALLOCATING/ERROR ({stuck})")
    return rep


def _n(db, sql, params=None) -> int:
    return db.q1(sql, params)["n"]


# ----------------------------------------------------------------- orchestration

def reset(db, settings) -> None:
    from urllib.parse import urlsplit
    host = urlsplit(db.url).hostname or "?"
    counts = db.table_counts()
    print(f"RESET: TRUNCATE schema '{settings.schema}' on host '{host}' "
          f"({sum(counts.values())} rows across {len(counts)} tables)")
    db.reset_all_data()
    print("  done – schema cleared.")


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(prog="stress_eval")
    ap.add_argument("--scenario", choices=["orders", "inbox", "both"], default="both")
    ap.add_argument("--states", help="comma list of state codes, e.g. KA,BR (default: all templates)")
    ap.add_argument("--lans", type=int, default=10, help="LANs per state (order sheet)")
    ap.add_argument("--papers", type=int, default=10, help="papers per state (uploaded)")
    ap.add_argument("--dup-upload", action="store_true", help="also upload an exact-copy duplicate per group")
    ap.add_argument("--max-cycles", type=int, default=5)
    ap.add_argument("--reset", action="store_true", help="TRUNCATE the schema before running")
    ap.add_argument("--live", action="store_true", help="run live (overrides dry_run) – issues real certificates")
    ap.add_argument("--config", default=str(ROOT / "config.yaml"))
    args = ap.parse_args(argv)

    settings = load_settings(args.config)
    if args.live:
        settings.raw["dry_run"] = False
    if settings.dry_run:
        print("dry_run is ON – pass --live to run a real e2e (nothing will be written otherwise). Aborting.")
        return 2   # abort before any DB/Google connection so a non-live run has zero side effects

    from estamp.db import DB
    from estamp.google_clients import Google
    db = DB(settings.env("db.url_env"), settings.schema, max_size=int(settings.get("workers", 8)) + 4)
    google = Google(settings)
    resolver = StateResolver(settings.states)

    try:
        states = [s.strip().upper() for s in args.states.split(",")] if args.states else None
        papers = load_papers(resolver)
        scenarios = ["inbox", "orders"] if args.scenario == "both" else [args.scenario]

        if args.reset:
            reset(db, settings)

        overall: list[tuple[str, Report]] = []
        for sc_name in scenarios:
            sc = build_scenario(settings, papers, states, args.lans, args.papers, sc_name, args.dup_upload)
            print(f"\n=== scenario: {sc_name} ===")
            print_plan(sc)
            print("Uploading assets:")
            upload_assets(google, settings, sc)
            print("Driving cycles:")
            ran = drive_cycles(settings, db, google, args.max_cycles)
            print(f"  drained after {ran} cycle(s)")
            rep = verify(db, settings, sc)
            overall.append((sc_name, rep))
            if args.scenario == "both" and sc_name == "inbox":
                # inbox papers are now stock; clear so the orders scenario starts clean
                reset(db, settings)

        return print_report(overall)
    finally:
        db.close()


def print_plan(sc: Scenario) -> None:
    print(f"{'state':<6}{'amt':>6}{'lans':>6}{'papers':>8}{'valid':>7}{'issuable':>10}"
          f"{'pending':>9}{'avail':>7}{'review':>7}  notes")
    for g in sc.groups:
        notes = "" if (g.active and g.single_amount) else ("inactive" if not g.active else "multi-amount→invalid")
        print(f"{g.state:<6}{g.amount:>6}{len(g.lans):>6}{len(g.papers):>8}{g.valid_lans:>7}"
              f"{g.issuable:>10}{g.pending:>9}{g.available:>7}{g.review:>7}  {notes}")


def print_report(overall: list[tuple[str, Report]]) -> int:
    print("\n" + "=" * 60 + "\nREPORT\n" + "=" * 60)
    all_ok = True
    for name, rep in overall:
        passed = sum(1 for ok, _ in rep.checks if ok)
        print(f"\n[{name}] {passed}/{len(rep.checks)} checks passed")
        for ok, msg in rep.checks:
            print(f"  {'✓' if ok else '✗'} {msg}")
        if rep.failures:
            all_ok = False
    print("\n" + ("ALL EXPECTATIONS MET ✓" if all_ok else "MISMATCHES FOUND ✗ – see ✗ lines above (what to fix)"))
    return 0 if all_ok else 1


if __name__ == "__main__":
    sys.exit(main())
